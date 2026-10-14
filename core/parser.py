from __future__ import annotations

import hashlib
import os
import zipfile
from pathlib import Path
from typing import Any, Dict, List, Optional, Tuple

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE, PP_PLACEHOLDER
from pptx.shapes.base import BaseShape

import re
from core.diff_engine import map_symbol_font_text, normalize_text
from core.models import (
    DeckMetadata,
    HyperlinkInfo,
    ParagraphInfo,
    ShapeInfo,
    SlideData,
    TableInfo,
)


def _extract_text_from_xml_element(element) -> str:
    """
    Extract text in document order from an XML element, including:
    - a:r (regular runs) with symbol font mapping
    - m:oMath, m:oMathPara, a14:m (OMML math) reading m:t
    - mc:AlternateContent (Choice then Fallback)
    - a:fld (fields)
    - a:br (line breaks)
    """
    if element is None:
        return ""
    parts: List[str] = []
    for child in element:
        tag = child.tag.split("}")[-1]
        if tag == "r":
            fn = None
            rPr = child.find("{http://schemas.openxmlformats.org/drawingml/2006/main}rPr")
            if rPr is not None:
                latin = rPr.find("{http://schemas.openxmlformats.org/drawingml/2006/main}latin")
                if latin is not None:
                    fn = latin.attrib.get("typeface")
            t_el = child.find("{http://schemas.openxmlformats.org/drawingml/2006/main}t")
            if t_el is not None and t_el.text:
                parts.append(map_symbol_font_text(t_el.text, fn))
        elif tag in ("m", "oMath", "oMathPara"):
            for sub in child.iter():
                subtag = sub.tag.split("}")[-1]
                if subtag in ("t",) and sub.text:
                    parts.append(map_symbol_font_text(sub.text))
        elif tag == "AlternateContent":
            choice_parts: List[str] = []
            for sub in child.iter():
                subtag = sub.tag.split("}")[-1]
                if subtag in ("Choice",):
                    for sub_c in sub.iter():
                        c_tag = sub_c.tag.split("}")[-1]
                        if c_tag in ("t",) and sub_c.text:
                            choice_parts.append(map_symbol_font_text(sub_c.text))
            if choice_parts:
                parts.extend(choice_parts)
            else:
                for sub in child.iter():
                    subtag = sub.tag.split("}")[-1]
                    if subtag in ("t",) and sub.text:
                        parts.append(map_symbol_font_text(sub.text))
        elif tag == "fld":
            t_el = child.find("{http://schemas.openxmlformats.org/drawingml/2006/main}t")
            if t_el is not None and t_el.text:
                parts.append(map_symbol_font_text(t_el.text))
        elif tag == "br":
            parts.append("\n")
        else:
            sub_res = _extract_text_from_xml_element(child)
            if sub_res:
                parts.append(sub_res)
    return "".join(parts)


def _extract_mtef_ole_text(part_blob: bytes) -> str:
    """
    Extract readable text from a MathType / Equation OLE object binary.
    """
    if not part_blob:
        return ""
    chars: List[str] = []
    i = 0
    length = len(part_blob)
    while i < length - 4:
        if part_blob[i] == 2:
            code = part_blob[i + 3] | (part_blob[i + 4] << 8)
            if (0x20 <= code <= 0x7E) or (0x390 <= code <= 0x3CF) or code == 0x2212 or code == 0x03C0:
                chars.append(chr(code))
                i += 5
                continue
        i += 1

    extracted = "".join(chars).strip()
    if extracted:
        extracted = re.sub(r"[\x00-\x1f\x7f-\x9f]", " ", extracted)
        extracted = re.sub(r"\s+", " ", extracted).strip()
        # De-duplicate identical repeated equation halves (common when both preview and native records exist)
        half = len(extracted) // 2
        if half > 5 and extracted[:half] == extracted[half:]:
            extracted = extracted[:half].strip()
        if len(re.findall(r"[A-Za-z0-9=+\-*/\(\)]", extracted)) >= 2:
            return extracted

    noise = {
        "mathtype", "equation", "dsmt4", "dsmt7", "winallbasiccodepages", "winallcodepages",
        "times new roman", "courier new", "mt extra", "appsmfcc", "design science, inc.",
        "system", "nani", "root entry", "ole", "equation native", "compobj", "olepres000",
    }
    raw_strings = re.findall(rb"[\x20-\x7e]{2,}", part_blob)
    cleaned: List[str] = []
    for s in raw_strings:
        t = s.decode("latin1", errors="ignore").strip()
        if not t or t.lower() in noise:
            continue
        if any(t.lower().startswith(n) for n in ("mathtype", "equation.", "dsmt", "apps", "winall")):
            continue
        if re.search(r"[A-Za-z0-9=+\-*/\(\)]", t):
            t_clean = re.sub(r"~([a-zA-Z0-9])", r"\1", t)
            t_clean = re.sub(r"\}", "", t_clean)
            cleaned.append(t_clean)
    return " ".join(cleaned)



class PresentationParseError(Exception):
    """Raised when a PPTX cannot be parsed or is invalid/corrupt."""
    pass


def validate_pptx_file(file_path: str | Path, file_label: str = "File") -> None:
    """
    Validate that the given path exists and is a valid PPTX file.
    Detects specific failure reasons and raises PresentationParseError with a friendly message:
    - File name starts with "~$": temporary lock file
    - File size 0 or under 1 KB: empty or incomplete upload
    - Magic bytes D0 CF 11 E0: legacy .ppt or password-protected file
    - Anything else: not a valid PPTX
    Identifies which file (original or remediated) encountered the issue.
    """
    path = Path(file_path)

    if not path.exists():
        raise PresentationParseError(f"{file_label} does not exist: {path.name}")

    # 1. File name starts with "~$"
    if path.name.startswith("~$"):
        raise PresentationParseError(
            f"{file_label} ('{path.name}'): This is a temporary lock file, not the real presentation."
        )

    # 2. File size 0 or under 1 KB (<= 1024 bytes)
    try:
        file_size = path.stat().st_size
    except Exception:
        file_size = 0

    if file_size <= 1024:
        raise PresentationParseError(
            f"{file_label} ('{path.name}'): File is empty or the upload was incomplete."
        )

    # 3. Magic bytes check (D0 CF 11 E0)
    try:
        with open(path, "rb") as f:
            magic = f.read(8)
    except Exception:
        magic = b""

    if len(magic) >= 4 and magic[:4] == b"\xd0\xcf\x11\xe0":
        raise PresentationParseError(
            f"{file_label} ('{path.name}'): This is an old .ppt or a password-protected file. Please re-save as .pptx."
        )

    # 4. Check extension
    if path.suffix.lower() not in [".pptx", ".potx", ".pptm"]:
        raise PresentationParseError(
            f"{file_label} ('{path.name}'): File is not a valid PPTX. Please open it in PowerPoint and Save As .pptx."
        )

    # 5. Check ZIP archive structure
    try:
        with zipfile.ZipFile(path, "r") as zf:
            names = zf.namelist()
            if not any(n.startswith("ppt/") for n in names):
                raise PresentationParseError(
                    f"{file_label} ('{path.name}'): File is not a valid PPTX. Please open it in PowerPoint and Save As .pptx."
                )
            for info in zf.infolist():
                if info.flag_bits & 0x1:
                    raise PresentationParseError(
                        f"{file_label} ('{path.name}'): This is an old .ppt or a password-protected file. Please re-save as .pptx."
                    )
    except PresentationParseError:
        raise
    except Exception:
        raise PresentationParseError(
            f"{file_label} ('{path.name}'): File is not a valid PPTX. Please open it in PowerPoint and Save As .pptx."
        )


def _get_paragraph_alignment(p) -> str:
    """Extract paragraph alignment name."""
    try:
        align = p.alignment
        if align is not None:
            name = getattr(align, "name", str(align))
            return name.replace("PP_ALIGN.", "")
    except Exception:
        pass
    return "DEFAULT (LEFT)"


def _get_vertical_anchor(tf) -> Optional[str]:
    """Extract text frame vertical anchor name."""
    try:
        anchor = tf.vertical_anchor
        if anchor is not None:
            name = getattr(anchor, "name", str(anchor))
            return name.replace("MSO_ANCHOR.", "")
    except Exception:
        pass
    return None


def _parse_table(shape: BaseShape, shape_id: int, shape_name: str) -> TableInfo:
    """Parse table structure and cell contents."""
    table = shape.table
    rows = len(table.rows)
    cols = len(table.columns)
    cell_matrix: List[List[str]] = []

    for row in table.rows:
        row_cells = []
        for cell in row.cells:
            cell_raw = cell.text or ""
            text = cell_raw.strip() if normalize_text(cell_raw) else ""
            row_cells.append(text)
        cell_matrix.append(row_cells)

    return TableInfo(
        shape_id=shape_id,
        name=shape_name,
        rows=rows,
        cols=cols,
        cell_texts=cell_matrix,
    )


def _extract_shape_data(
    shape: BaseShape,
    z_index: int,
    text_blocks: List[str],
    slide: Optional[Any] = None,
    resolve_inherited_fonts: bool = False,
) -> ShapeInfo:
    """Extract positioning in EMU, alignment, rotation, text, font styles, and media hashes."""
    shape_id = getattr(shape, "shape_id", z_index)
    name = shape.name or f"Shape {z_index}"

    try:
        left_emu = int(shape.left) if shape.left is not None else 0
        top_emu = int(shape.top) if shape.top is not None else 0
        width_emu = int(shape.width) if shape.width is not None else 0
        height_emu = int(shape.height) if shape.height is not None else 0
    except Exception:
        left_emu, top_emu, width_emu, height_emu = 0, 0, 0, 0

    # Resolve inherited positions from layout and master placeholders if shape has no valid position
    if (width_emu <= 0 or height_emu <= 0 or (left_emu == 0 and top_emu == 0 and width_emu == 0 and height_emu == 0)) and getattr(shape, "is_placeholder", False):
        try:
            ph_idx = getattr(shape.placeholder_format, "idx", None)
            ph_type = getattr(shape.placeholder_format, "type", None)
            layout = getattr(slide, "slide_layout", None)
            master = getattr(layout, "slide_master", None) if layout else None

            # 1. Search slide layout placeholders
            resolved = False
            if layout and hasattr(layout, "placeholders"):
                # First try matching by idx
                for l_ph in layout.placeholders:
                    l_fmt = getattr(l_ph, "placeholder_format", None)
                    if l_fmt and ph_idx is not None and l_fmt.idx == ph_idx:
                        if getattr(l_ph, "width", 0) and getattr(l_ph, "height", 0):
                            left_emu = int(l_ph.left) if l_ph.left is not None else left_emu
                            top_emu = int(l_ph.top) if l_ph.top is not None else top_emu
                            width_emu = int(l_ph.width) if l_ph.width is not None else width_emu
                            height_emu = int(l_ph.height) if l_ph.height is not None else height_emu
                            resolved = True
                            break
                # If not matched by idx, try matching by type
                if not resolved and ph_type is not None:
                    for l_ph in layout.placeholders:
                        l_fmt = getattr(l_ph, "placeholder_format", None)
                        if l_fmt and l_fmt.type == ph_type:
                            if getattr(l_ph, "width", 0) and getattr(l_ph, "height", 0):
                                left_emu = int(l_ph.left) if l_ph.left is not None else left_emu
                                top_emu = int(l_ph.top) if l_ph.top is not None else top_emu
                                width_emu = int(l_ph.width) if l_ph.width is not None else width_emu
                                height_emu = int(l_ph.height) if l_ph.height is not None else height_emu
                                resolved = True
                                break

            # 2. Search master placeholders if still unresolved
            if not resolved and master and hasattr(master, "placeholders"):
                if ph_type is not None:
                    for m_ph in master.placeholders:
                        m_fmt = getattr(m_ph, "placeholder_format", None)
                        if m_fmt and m_fmt.type == ph_type:
                            if getattr(m_ph, "width", 0) and getattr(m_ph, "height", 0):
                                left_emu = int(m_ph.left) if m_ph.left is not None else left_emu
                                top_emu = int(m_ph.top) if m_ph.top is not None else top_emu
                                width_emu = int(m_ph.width) if m_ph.width is not None else width_emu
                                height_emu = int(m_ph.height) if m_ph.height is not None else height_emu
                                resolved = True
                                break
        except Exception:
            pass

    rotation = float(getattr(shape, "rotation", 0.0) or 0.0)
    shape_type_str = str(shape.shape_type)

    is_image = False
    image_hash = None
    is_chart = False
    is_table = False
    is_group = False
    is_placeholder = False
    placeholder_type = None

    if shape.is_placeholder:
        is_placeholder = True
        try:
            placeholder_type = str(shape.placeholder_format.type)
        except Exception:
            placeholder_type = "UNKNOWN"

    if shape.shape_type == MSO_SHAPE_TYPE.PICTURE or shape.element.tag.endswith("pic"):
        is_image = True
        try:
            if hasattr(shape, "image") and shape.image and shape.image.blob:
                image_hash = hashlib.sha256(shape.image.blob).hexdigest()
        except Exception:
            image_hash = None
    elif shape.shape_type == MSO_SHAPE_TYPE.CHART or getattr(shape, "has_chart", False):
        is_chart = True
    elif shape.shape_type == MSO_SHAPE_TYPE.TABLE or getattr(shape, "has_table", False):
        is_table = True
    elif shape.shape_type == MSO_SHAPE_TYPE.GROUP:
        is_group = True

    # Text & Font extraction
    text = ""
    paragraphs_info: List[ParagraphInfo] = []
    vertical_anchor = None
    hyperlinks: List[HyperlinkInfo] = []
    primary_font_name = None
    primary_font_size_pt = None
    primary_is_bold = None
    primary_is_italic = None
    primary_is_underline = None
    primary_color_hex = None

    if getattr(shape, "has_text_frame", False):
        tf = shape.text_frame
        vertical_anchor = _get_vertical_anchor(tf)

        para_texts: List[str] = []
        for p in tf.paragraphs:
            p_elem = getattr(p, "_p", None)
            p_xml_text = _extract_text_from_xml_element(p_elem).strip() if p_elem is not None else ""
            p_text = p_xml_text if p_xml_text else (p.text or "").strip()
            align_name = _get_paragraph_alignment(p)

            p_font_name = None
            p_font_size_pt = None
            p_is_bold = None
            p_is_italic = None
            p_is_underline = None
            p_color_hex = None

            for r in p.runs:
                r_text = (r.text or "").strip()
                if r_text and normalize_text(r_text):
                    try:
                        if p_font_name is None and r.font.name:
                            p_font_name = r.font.name
                        if p_font_size_pt is None and r.font.size:
                            p_font_size_pt = round(float(r.font.size.pt), 1)
                        if p_is_bold is None and r.font.bold is not None:
                            p_is_bold = r.font.bold
                        if p_is_italic is None and r.font.italic is not None:
                            p_is_italic = r.font.italic
                        if p_is_underline is None and r.font.underline is not None:
                            p_is_underline = r.font.underline
                        if p_color_hex is None and r.font.color and hasattr(r.font.color, "rgb") and r.font.color.rgb:
                            p_color_hex = str(r.font.color.rgb)
                    except Exception:
                        pass

                try:
                    if r.hyperlink and r.hyperlink.address:
                        hyperlinks.append(
                            HyperlinkInfo(
                                text=(r.text or "").strip(),
                                url=r.hyperlink.address.strip(),
                                shape_name=name,
                            )
                        )
                except Exception:
                    pass

            if p_text:
                para_texts.append(p_text)

            paragraphs_info.append(
                ParagraphInfo(
                    text=p_text,
                    alignment=align_name,
                    font_name=p_font_name,
                    font_size_pt=p_font_size_pt,
                    is_bold=p_is_bold,
                    is_italic=p_is_italic,
                    is_underline=p_is_underline,
                    color_hex=p_color_hex,
                )
            )

            if primary_font_name is None and p_font_name:
                primary_font_name = p_font_name
            if primary_font_size_pt is None and p_font_size_pt:
                primary_font_size_pt = p_font_size_pt
            if primary_is_bold is None and p_is_bold is not None:
                primary_is_bold = p_is_bold
            if primary_is_italic is None and p_is_italic is not None:
                primary_is_italic = p_is_italic
            if primary_is_underline is None and p_is_underline is not None:
                primary_is_underline = p_is_underline
            if primary_color_hex is None and p_color_hex is not None:
                primary_color_hex = p_color_hex

        if para_texts:
            text = "\n".join(para_texts).strip()
            if normalize_text(text):
                text_blocks.append(text)

        # Resolve inherited fonts from layout placeholder if check_fonts enabled
        if resolve_inherited_fonts and is_placeholder and slide and hasattr(slide, "slide_layout"):
            try:
                ph_idx = shape.placeholder_format.idx
                for l_ph in slide.slide_layout.placeholders:
                    if getattr(l_ph, "placeholder_format", None) and l_ph.placeholder_format.idx == ph_idx:
                        if getattr(l_ph, "has_text_frame", False):
                            for lp in l_ph.text_frame.paragraphs:
                                if primary_font_name is None and hasattr(lp, "font") and lp.font.name:
                                    primary_font_name = lp.font.name
                                if primary_font_size_pt is None and hasattr(lp, "font") and lp.font.size:
                                    primary_font_size_pt = round(float(lp.font.size.pt), 1)
                                for lr in lp.runs:
                                    if primary_font_name is None and lr.font.name:
                                        primary_font_name = lr.font.name
                                    if primary_font_size_pt is None and lr.font.size:
                                        primary_font_size_pt = round(float(lr.font.size.pt), 1)
            except Exception:
                pass

    # Check for OLE or OMML equations in shapes without text frames
    if not text:
        elem_text = _extract_text_from_xml_element(shape.element).strip()
        if elem_text and normalize_text(elem_text):
            text = elem_text
            text_blocks.append(text)

        # Check for oleObj embeddings
        seen_ole_rids: Set[str] = set()
        ole_els = shape.element.findall(".//{http://schemas.openxmlformats.org/presentationml/2006/main}oleObj")
        for ole_el in ole_els:
            r_id = ole_el.attrib.get("{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id")
            if r_id and r_id not in seen_ole_rids and slide and hasattr(slide, "part") and r_id in slide.part.rels:
                seen_ole_rids.add(r_id)
                try:
                    rel = slide.part.rels[r_id]
                    ole_txt = _extract_mtef_ole_text(rel.target_part.blob)
                    if ole_txt and normalize_text(ole_txt):
                        text = (text + " " + ole_txt).strip() if text else ole_txt
                        text_blocks.append(text)
                except Exception:
                    pass

    # Check for fallback picture inside oleObj or AlternateContent
    if not is_image and not image_hash:
        pic_els = shape.element.findall(".//{http://schemas.openxmlformats.org/presentationml/2006/main}pic")
        for pic_el in pic_els:
            blip = pic_el.find(".//{http://schemas.openxmlformats.org/drawingml/2006/main}blip")
            if blip is not None:
                r_id = blip.attrib.get("{http://schemas.openxmlformats.org/officeDocument/2006/relationships}embed")
                if r_id and slide and hasattr(slide, "part") and r_id in slide.part.rels:
                    try:
                        rel = slide.part.rels[r_id]
                        blob = rel.target_part.blob
                        if blob:
                            image_hash = hashlib.sha256(blob).hexdigest()
                            is_image = True
                            break
                    except Exception:
                        pass

    # Shape click action hyperlink
    try:
        if shape.click_action and shape.click_action.hyperlink and shape.click_action.hyperlink.address:
            hyperlinks.append(
                HyperlinkInfo(
                    text=text or name,
                    url=shape.click_action.hyperlink.address.strip(),
                    shape_name=name,
                )
            )
    except Exception:
        pass

    # Recursively unpack group shapes for text and images
    if is_group and hasattr(shape, "shapes"):
        for sub_shape in shape.shapes:
            try:
                sub_txt = ""
                if getattr(sub_shape, "has_text_frame", False) and sub_shape.text_frame:
                    p_xml = [_extract_text_from_xml_element(getattr(p, "_p", None)).strip() for p in sub_shape.text_frame.paragraphs]
                    sub_txt = "\n".join([pt for pt in p_xml if pt]).strip() or (sub_shape.text_frame.text or "").strip()
                else:
                    sub_txt = _extract_text_from_xml_element(sub_shape.element).strip()

                if sub_txt and normalize_text(sub_txt):
                    text = (text + "\n" + sub_txt).strip() if text else sub_txt
                    text_blocks.append(sub_txt)

                if (sub_shape.shape_type == MSO_SHAPE_TYPE.PICTURE or sub_shape.element.tag.endswith("pic")) and hasattr(sub_shape, "image"):
                    if not image_hash and sub_shape.image and sub_shape.image.blob:
                        image_hash = hashlib.sha256(sub_shape.image.blob).hexdigest()
            except Exception:
                pass

    return ShapeInfo(
        shape_id=shape_id,
        name=name,
        shape_type=shape_type_str,
        left_emu=left_emu,
        top_emu=top_emu,
        width_emu=width_emu,
        height_emu=height_emu,
        rotation=rotation,
        z_order=z_index,
        text=text,
        paragraphs=paragraphs_info,
        vertical_anchor=vertical_anchor,
        font_name=primary_font_name,
        font_size_pt=primary_font_size_pt,
        is_bold=primary_is_bold,
        is_italic=primary_is_italic,
        is_underline=primary_is_underline,
        color_hex=primary_color_hex,
        is_image=is_image,
        image_hash=image_hash,
        is_chart=is_chart,
        is_table=is_table,
        is_group=is_group,
        is_placeholder=is_placeholder,
        placeholder_type=placeholder_type,
        hyperlinks=hyperlinks,
    )


def parse_presentation(
    file_path: str | Path,
    file_label: str = "Presentation file",
    resolve_inherited_fonts: bool = False,
) -> Tuple[DeckMetadata, List[SlideData]]:
    """
    Parse a PowerPoint presentation into metadata and slide structures.
    """
    validate_pptx_file(file_path, file_label=file_label)

    try:
        prs = Presentation(str(file_path))
    except PresentationParseError:
        raise
    except Exception as e:
        raise PresentationParseError(
            f"{file_label} ('{Path(file_path).name}'): File is not a valid PPTX. Please open it in PowerPoint and Save As .pptx."
        ) from e

    slide_w = int(getattr(prs, "slide_width", 9144000) or 9144000)
    slide_h = int(getattr(prs, "slide_height", 6858000) or 6858000)

    metadata = DeckMetadata(
        title=prs.core_properties.title or "",
        author=prs.core_properties.author or "",
        slide_count=len(prs.slides),
        slide_width_emu=slide_w,
        slide_height_emu=slide_h,
    )

    slides_data: List[SlideData] = []

    for idx, slide in enumerate(prs.slides, start=1):
        slide_id = getattr(slide, "slide_id", idx)

        # Title
        title_text = ""
        try:
            if slide.shapes.title and slide.shapes.title.text:
                title_text = slide.shapes.title.text.strip()
        except Exception:
            pass

        slide_text_blocks: List[str] = []
        shapes_info: List[ShapeInfo] = []
        tables_info: List[TableInfo] = []
        slide_hyperlinks: List[HyperlinkInfo] = []
        img_count = 0
        chart_count = 0
        table_count = 0

        for z_idx, shape in enumerate(slide.shapes):
            shape_info = _extract_shape_data(
                shape, z_idx, slide_text_blocks, slide=slide, resolve_inherited_fonts=resolve_inherited_fonts
            )
            shapes_info.append(shape_info)
            slide_hyperlinks.extend(shape_info.hyperlinks)

            if shape_info.is_image:
                img_count += 1
            elif shape_info.is_chart:
                chart_count += 1
            elif shape_info.is_table:
                table_count += 1
                table_info = _parse_table(shape, shape_info.shape_id, shape_info.name)
                tables_info.append(table_info)
                # Add table cell texts to slide text blocks
                for row_cells in table_info.cell_texts:
                    for cell_txt in row_cells:
                        if normalize_text(cell_txt):
                            slide_text_blocks.append(cell_txt)

        # Slide notes
        notes_text = ""
        try:
            if slide.has_notes_slide and slide.notes_slide.notes_text_frame:
                raw_notes = slide.notes_slide.notes_text_frame.text or ""
                notes_text = raw_notes.strip() if normalize_text(raw_notes) else ""
        except Exception:
            notes_text = ""

        full_text = "\n".join(slide_text_blocks)

        slide_data = SlideData(
            slide_id=slide_id,
            slide_index=idx,
            title=title_text,
            text_blocks=slide_text_blocks,
            text_content=full_text,
            shapes=shapes_info,
            tables=tables_info,
            notes=notes_text,
            hyperlinks=slide_hyperlinks,
            images_count=img_count,
            charts_count=chart_count,
            tables_count=table_count,
            shapes_count=len(shapes_info),
            slide_width_emu=slide_w,
            slide_height_emu=slide_h,
        )
        slides_data.append(slide_data)

    return metadata, slides_data
