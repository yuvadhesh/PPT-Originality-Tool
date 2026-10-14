from __future__ import annotations

import openpyxl
from pathlib import Path
from PIL import Image

from core.models import (
    BoundingBox,
    DeckComparisonResult,
    DeckMetadata,
    Issue,
    RuleType,
    SlideDiffResult,
    SlideMatch,
    SlideStatus,
)
from core.report import export_excel_report, export_html_report
from core.screenshot_annotator import (
    apply_dimming_overlay,
    create_side_by_side_highlight_image,
    emu_to_pixel,
)


def _create_mock_comparison_result() -> DeckComparisonResult:
    iss1 = Issue(
        slide_num=1,
        element_name="Bullet 1",
        rule_type=RuleType.CONTENT,
        issue_type="Missing Text",
        description="Paragraph from original is missing.",
        original_value="Revenue exceeded $10M.",
        remediated_value="Missing",
        difference="Not found",
        is_error=True,
        issue_num=1,
        orig_box=BoundingBox(left_emu=914400, top_emu=914400, width_emu=1828800, height_emu=914400),
    )
    iss2 = Issue(
        slide_num=2,
        element_name="TargetBox",
        rule_type=RuleType.ALIGNMENT,
        issue_type="Shifted Position (Left)",
        description="Horizontal position shifted beyond tolerance.",
        original_value="914400 EMU",
        remediated_value="1828800 EMU",
        difference="+914400 EMU",
        is_error=True,
        issue_num=2,
        orig_box=BoundingBox(left_emu=914400, top_emu=914400, width_emu=1828800, height_emu=914400),
        rem_box=BoundingBox(left_emu=1828800, top_emu=914400, width_emu=1828800, height_emu=914400),
    )
    iss3 = Issue(
        slide_num=2,
        element_name="TitleBox",
        rule_type=RuleType.FONT,
        issue_type="Changed Font Size",
        description="Font size changed from 24pt to 16pt.",
        original_value="24.0 pt",
        remediated_value="16.0 pt",
        difference="-8.0 pt",
        is_error=True,
        issue_num=3,
        orig_box=BoundingBox(left_emu=914400, top_emu=457200, width_emu=3657600, height_emu=914400),
        rem_box=BoundingBox(left_emu=914400, top_emu=457200, width_emu=3657600, height_emu=914400),
    )

    s1 = SlideDiffResult(
        slide_key="Slide 1",
        orig_index=1,
        rem_index=1,
        match_type="MATCHED",
        similarity_score=90.0,
        issues=[iss1],
        status=SlideStatus.FAIL,
    )
    s2 = SlideDiffResult(
        slide_key="Slide 2",
        orig_index=2,
        rem_index=2,
        match_type="MATCHED",
        similarity_score=95.0,
        issues=[iss2, iss3],
        status=SlideStatus.FAIL,
    )

    return DeckComparisonResult(
        orig_filename="original.pptx",
        rem_filename="remediated.pptx",
        orig_metadata=DeckMetadata(title="Original", slide_count=2, slide_width_emu=9144000, slide_height_emu=6858000),
        rem_metadata=DeckMetadata(title="Remediated", slide_count=2, slide_width_emu=9144000, slide_height_emu=6858000),
        slide_matches=[
            SlideMatch(orig_index=1, rem_index=1, match_type="MATCHED", similarity_score=90.0),
            SlideMatch(orig_index=2, rem_index=2, match_type="MATCHED", similarity_score=95.0),
        ],
        slide_results=[s1, s2],
        deck_issues=[],
        processing_time_sec=1.0,
    )


def test_emu_to_pixel_conversion():
    # 1 inch = 914400 EMU. Slide width = 10 inches = 9144000 EMU. Image width = 1000 px.
    # Therefore 1 inch in EMU should equal 100 pixels.
    px_x, px_y, px_w, px_h = emu_to_pixel(
        emu_x=914400,
        emu_y=1828800,
        emu_w=2743200,
        emu_h=914400,
        slide_w_emu=9144000,
        slide_h_emu=6858000,
        img_w_px=1000,
        img_h_px=750,
    )
    assert px_x == 100
    assert px_y == 200
    assert px_w == 300
    assert px_h == 100


def test_bounding_box_drawing_and_overlay(tmp_path):
    img = Image.new("RGB", (800, 600), (200, 200, 200))
    boxes = [(100, 100, 200, 150)]

    # Test dimming overlay
    dimmed = apply_dimming_overlay(img, boxes, padding_px=20, dim_percent=0.30)
    assert dimmed.size == (800, 600)

    # Test side-by-side creation
    res = _create_mock_comparison_result()
    combined = create_side_by_side_highlight_image(
        orig_img=img,
        rem_img=img,
        issues=res.slide_results[1].issues,
        slide_w_emu=9144000,
        slide_h_emu=6858000,
        max_width_px=1200,
    )
    assert combined.width <= 1200
    assert combined.height > 300


def test_export_excel_report_with_screenshots_and_hyperlinks(tmp_path):
    res = _create_mock_comparison_result()
    excel_path = tmp_path / "test_report.xlsx"

    # Create dummy screenshot images for testing embedding
    shot1 = tmp_path / "shot1.png"
    shot2 = tmp_path / "shot2.png"
    img = Image.new("RGB", (600, 400), (230, 230, 230))
    img.save(str(shot1))
    img.save(str(shot2))

    screenshots_map = {1: shot1, 2: shot2}

    export_excel_report(res, excel_path, screenshots_map=screenshots_map)

    assert excel_path.exists()
    wb = openpyxl.load_workbook(str(excel_path))
    assert "Summary" in wb.sheetnames
    assert "Issues" in wb.sheetnames
    assert "Screenshots" in wb.sheetnames

    ws_issues = wb["Issues"]
    # Check column headers
    assert ws_issues.cell(row=1, column=1).value == "Issue #"
    assert ws_issues.cell(row=1, column=3).value == "Rule"
    assert ws_issues.cell(row=1, column=10).value == "Screenshot Ref"

    # Check hyperlinks exist
    ref_cell = ws_issues.cell(row=2, column=10)
    assert ref_cell.hyperlink is not None
    assert "Screenshots" in ref_cell.hyperlink.target

    ws_shots = wb["Screenshots"]
    assert len(ws_shots._images) >= 1

    ws_summary = wb["Summary"]
    assert ws_summary["A1"].value == "PPT Content, Alignment & Font QA Report"


def test_export_html_report(tmp_path):
    res = _create_mock_comparison_result()
    html_path = tmp_path / "test_report.html"
    export_html_report(res, html_path)

    assert html_path.exists()
    content = html_path.read_text(encoding="utf-8")
    assert "Revenue exceeded $10M" in content
    assert "Checks Run: <strong>Content, Alignment</strong>" in content
    assert "Font Issues" not in content


def test_export_reports_with_check_fonts_toggle(tmp_path):
    res_false = _create_mock_comparison_result()
    res_false.check_fonts = False

    excel_false = tmp_path / "report_false.xlsx"
    export_excel_report(res_false, excel_false)

    wb_false = openpyxl.load_workbook(str(excel_false))
    ws_sum_false = wb_false["Summary"]
    summary_text_false = [ws_sum_false.cell(row=r, column=1).value for r in range(1, 15)]
    assert any(ws_sum_false.cell(row=r, column=2).value == "Content, Alignment" for r in range(1, 15))
    assert "Total Font Issues" not in summary_text_false

    ws_iss_false = wb_false["Issues"]
    # Header is row 1. Issues follow. None should be Font
    rule_values_false = [ws_iss_false.cell(row=r, column=3).value for r in range(2, ws_iss_false.max_row + 1)]
    assert "Font" not in rule_values_false

    # Now test with check_fonts = True
    res_true = _create_mock_comparison_result()
    res_true.check_fonts = True

    excel_true = tmp_path / "report_true.xlsx"
    export_excel_report(res_true, excel_true)

    wb_true = openpyxl.load_workbook(str(excel_true))
    ws_sum_true = wb_true["Summary"]
    summary_text_true = [ws_sum_true.cell(row=r, column=1).value for r in range(1, 15)]
    assert any(ws_sum_true.cell(row=r, column=2).value == "Content, Alignment, Font" for r in range(1, 15))
    assert "Total Font Issues" in summary_text_true

    ws_iss_true = wb_true["Issues"]
    rule_values_true = [ws_iss_true.cell(row=r, column=3).value for r in range(2, ws_iss_true.max_row + 1)]
    assert "Font" in rule_values_true


def test_passing_slide_excluded_from_screenshots_sheet(tmp_path):
    """Verify that slides with status Pass are strictly excluded from the Screenshots sheet."""
    res = _create_mock_comparison_result()
    # Mark Slide 1 as PASS
    res.slide_results[0].status = SlideStatus.PASS
    # Keep Slide 2 as FAIL

    excel_path = tmp_path / "test_pass_excluded.xlsx"
    shot1 = tmp_path / "shot1.png"
    shot2 = tmp_path / "shot2.png"
    img = Image.new("RGB", (600, 400), (230, 230, 230))
    img.save(str(shot1))
    img.save(str(shot2))
    screenshots_map = {1: shot1, 2: shot2}

    export_excel_report(res, excel_path, screenshots_map=screenshots_map)

    wb = openpyxl.load_workbook(str(excel_path))
    ws_shots = wb["Screenshots"]

    # Gather all cell values from Screenshots sheet
    shot_cell_values = [
        str(cell.value) for row in ws_shots.iter_rows(values_only=False) for cell in row if cell.value
    ]

    # Slide 1 (Passing) MUST NOT appear in the Screenshots sheet
    assert not any("Slide 1" in v for v in shot_cell_values), "Slide 1 is passing and must not appear in Screenshots"
    assert not any("Pass" in v for v in shot_cell_values), "The word 'Pass' must never appear in Screenshots sheet"

    # Slide 2 (Failing) MUST appear with correct format: "Slide N - Issues: #..."
    assert any("Slide 2 - Issues:" in v for v in shot_cell_values), "Failing slide 2 must appear with 'Slide 2 - Issues: ...' header"
    # Ensure (#1) style label does not appear
    assert not any("(#" in v for v in shot_cell_values), "Legacy (#1) style heading label must be removed"
