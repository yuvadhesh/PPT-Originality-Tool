from __future__ import annotations

import os
from pathlib import Path
from typing import Tuple

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt
from PIL import Image, ImageDraw


def create_sample_image(output_path: str | Path, color=(70, 130, 180), text="Sample Graphic") -> Path:
    """Create a sample PNG image to embed into test presentations."""
    path = Path(output_path)
    path.parent.mkdir(parents=True, exist_ok=True)
    img = Image.new("RGB", (300, 200), color=color)
    draw = ImageDraw.Draw(img)
    draw.rectangle([10, 10, 290, 190], outline=(255, 255, 255), width=3)
    img.save(str(path))
    return path


def create_sample_decks(output_dir: str | Path) -> Tuple[Path, Path]:
    """
    Generate an ORIGINAL and a REMEDIATED .pptx presentation pair covering:
    - Missing text (Rule 1)
    - Missing image (Rule 1)
    - Shifted shape position (Rule 2)
    - Resized shape width/height (Rule 2)
    - Changed text paragraph alignment (Rule 2)
    """
    out_dir = Path(output_dir)
    out_dir.mkdir(parents=True, exist_ok=True)

    img1_path = create_sample_image(out_dir / "img1.png", (41, 128, 185), "Img 1")
    img2_path = create_sample_image(out_dir / "img2.png", (39, 174, 96), "Img 2")

    orig_path = out_dir / "sample_original.pptx"
    rem_path = out_dir / "sample_remediated.pptx"

    # =========================================================================
    # 1. ORIGINAL PRESENTATION
    # =========================================================================
    prs_orig = Presentation()
    prs_orig.core_properties.title = "Sample Deck - Original"

    # Slide 1: Clean Matching Slide (Pass)
    s1 = prs_orig.slides.add_slide(prs_orig.slide_layouts[0])
    s1.shapes.title.text = "Introduction"
    s1.placeholders[1].text = "Welcome to the presentation."

    # Slide 2: Text that will be missing in remediated (Rule 1)
    s2 = prs_orig.slides.add_slide(prs_orig.slide_layouts[1])
    s2.shapes.title.text = "Financial Results"
    tf2 = s2.shapes.placeholders[1].text_frame
    tf2.text = "Q3 financial highlights:"
    p2 = tf2.add_paragraph()
    p2.text = "Revenue grew by 24% year-over-year reaching 15 million dollars across all territories."
    p3 = tf2.add_paragraph()
    p3.text = "Customer retention maintained at 98%."

    # Slide 3: Images (one will be missing in remediated) (Rule 1)
    s3 = prs_orig.slides.add_slide(prs_orig.slide_layouts[5])
    s3.shapes.title.text = "Visual Assets"
    s3.shapes.add_picture(str(img1_path), Inches(1), Inches(2), Inches(3), Inches(2))
    s3.shapes.add_picture(str(img2_path), Inches(5), Inches(2), Inches(3), Inches(2))

    # Slide 4: Shape that will be shifted in remediated (Rule 2)
    s4 = prs_orig.slides.add_slide(prs_orig.slide_layouts[5])
    s4.shapes.title.text = "Position Alignment Slide"
    box_s4 = s4.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(2), Inches(2), Inches(3), Inches(1.5))
    box_s4.name = "TargetBox"
    box_s4.text_frame.text = "Static Shape Content"

    # Slide 5: Shape that will be resized in remediated (Rule 2)
    s5 = prs_orig.slides.add_slide(prs_orig.slide_layouts[5])
    s5.shapes.title.text = "Dimension Check Slide"
    box_s5 = s5.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(2), Inches(2), Inches(4), Inches(2))
    box_s5.name = "ResizeBox"
    box_s5.text_frame.text = "Dimension Check Box"

    # Slide 6: Paragraph text alignment will be changed in remediated (Rule 2)
    s6 = prs_orig.slides.add_slide(prs_orig.slide_layouts[5])
    s6.shapes.title.text = "Typography Alignment Slide"
    tb_s6 = s6.shapes.add_textbox(Inches(2), Inches(2), Inches(5), Inches(2))
    tb_s6.name = "AlignedTextBox"
    p_s6 = tb_s6.text_frame.paragraphs[0]
    p_s6.text = "Centered Important Statement"
    p_s6.alignment = PP_ALIGN.CENTER  # Originally CENTER

    # Slide 7: Text that will have font change in remediated (Arial -> Times New Roman)
    s7 = prs_orig.slides.add_slide(prs_orig.slide_layouts[5])
    s7.shapes.title.text = "Font Preservation Slide"
    tb_s7 = s7.shapes.add_textbox(Inches(2), Inches(2), Inches(5), Inches(2))
    tb_s7.name = "FontTextBox"
    p_s7 = tb_s7.text_frame.paragraphs[0]
    p_s7.text = "Heading Text in Arial 24pt"
    r_s7 = p_s7.runs[0]
    r_s7.font.name = "Arial"
    r_s7.font.size = Pt(24)

    prs_orig.save(str(orig_path))

    # =========================================================================
    # 2. REMEDIATED PRESENTATION (with specific Rule 1, Rule 2, and Font defects)
    # =========================================================================
    prs_rem = Presentation()
    prs_rem.core_properties.title = "Sample Deck - Remediated"

    # Slide 1: Identical match (Pass)
    sr1 = prs_rem.slides.add_slide(prs_rem.slide_layouts[0])
    sr1.shapes.title.text = "Introduction"
    sr1.placeholders[1].text = "Welcome to the presentation."

    # Slide 2: Missing text ("Revenue grew by 24%...") (Rule 1 error)
    sr2 = prs_rem.slides.add_slide(prs_rem.slide_layouts[1])
    sr2.shapes.title.text = "Financial Results"
    tfr2 = sr2.shapes.placeholders[1].text_frame
    tfr2.text = "Q3 financial highlights:"
    # Omitted paragraph 2!
    pr3 = tfr2.add_paragraph()
    pr3.text = "Customer retention maintained at 98%."

    # Slide 3: Missing image (img2 omitted) (Rule 1 error)
    sr3 = prs_rem.slides.add_slide(prs_rem.slide_layouts[5])
    sr3.shapes.title.text = "Visual Assets"
    sr3.shapes.add_picture(str(img1_path), Inches(1), Inches(2), Inches(3), Inches(2))
    # img2 is missing!

    # Slide 4: Shifted shape (Left shifted by +0.75 inch = 685800 EMU > 45720 EMU) (Rule 2 error)
    sr4 = prs_rem.slides.add_slide(prs_rem.slide_layouts[5])
    sr4.shapes.title.text = "Position Alignment Slide"
    box_sr4 = sr4.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(2.75), Inches(2), Inches(3), Inches(1.5))
    box_sr4.name = "TargetBox"
    box_sr4.text_frame.text = "Static Shape Content"

    # Slide 5: Resized shape (Width shrunk from 4.0 inches to 2.5 inches) (Rule 2 error)
    sr5 = prs_rem.slides.add_slide(prs_rem.slide_layouts[5])
    sr5.shapes.title.text = "Dimension Check Slide"
    box_sr5 = sr5.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(2), Inches(2), Inches(2.5), Inches(2))
    box_sr5.name = "ResizeBox"
    box_sr5.text_frame.text = "Dimension Check Box"

    # Slide 6: Changed text alignment (Changed from CENTER to LEFT) (Rule 2 error)
    sr6 = prs_rem.slides.add_slide(prs_rem.slide_layouts[5])
    sr6.shapes.title.text = "Typography Alignment Slide"
    tb_sr6 = sr6.shapes.add_textbox(Inches(2), Inches(2), Inches(5), Inches(2))
    tb_sr6.name = "AlignedTextBox"
    p_sr6 = tb_sr6.text_frame.paragraphs[0]
    p_sr6.text = "Centered Important Statement"
    p_sr6.alignment = PP_ALIGN.LEFT  # Changed to LEFT!

    # Slide 7: Changed font name (Arial -> Times New Roman) (Font Rule error)
    sr7 = prs_rem.slides.add_slide(prs_rem.slide_layouts[5])
    sr7.shapes.title.text = "Font Preservation Slide"
    tb_sr7 = sr7.shapes.add_textbox(Inches(2), Inches(2), Inches(5), Inches(2))
    tb_sr7.name = "FontTextBox"
    p_sr7 = tb_sr7.text_frame.paragraphs[0]
    p_sr7.text = "Heading Text in Arial 24pt"
    r_sr7 = p_sr7.runs[0]
    r_sr7.font.name = "Times New Roman"  # Changed from Arial to Times New Roman!
    r_sr7.font.size = Pt(24)

    prs_rem.save(str(rem_path))

    return orig_path, rem_path


if __name__ == "__main__":
    o, r = create_sample_decks("tests/sample_decks")
    print(f"Sample decks created: {o}, {r}")
